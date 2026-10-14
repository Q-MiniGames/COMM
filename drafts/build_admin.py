"""Prototype: RevRecap + Transaction History + Rental Data Tracker -> Admin commission sheet.

Usage: python3 -I build_admin.py REVRECAP TRANSHIST TRACKER TEMPLATE_ADMIN OUT.xlsx
"""
import re
import sys
from copy import copy

import openpyxl
import pandas as pd
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment

REVRECAP, TRANSHIST, TRACKER, TEMPLATE, OUT = sys.argv[1:6]

# ---------------------------------------------------------------- lookup tables
# FSP rate per chargeable hour, keyed by admin model label.  Inferred from the
# SEP-2026 admin (FSP / hours); "?" ones are low confidence and need confirming.
FSP_RATE = {
    "320GX": 65, "320GX + H": 75, "DP50CNT": 35, "SEM922F": 75,
    "426F2LRC + H": 37, "426F2 + H": 37, "336-XZPI": 100,
}
# Comm. rate by model, as used in SEP-2026 admin.
COMM_RATE = {
    "320GX": .01, "320GX + H": .01, "340D2CXL + H": .01, "966GC": .01,
    "336-XZPI": .01, "SEM922F": .01, "TH": .01,
    "216B3WKLR": .015, "226B3WKLR": .015, "426F2LRC + H": .015, "426F2 + H": .015,
    "303.5": .015, "DP50CNT": .0175,
}
# Fleet-deal factor and new-customer flag by customer no., from SEP-2026 admin.
FLEET = {"QP45280": .7, "QE01110": .8, "QN00785": .9}
NEW_CUSTOMER = {"QP53540", "QP50760"}

YELLOW = PatternFill("solid", fgColor="FFFF00")
ORANGE = PatternFill("solid", fgColor="FFC000")


def clean(s):
    """Strip Excel CR escapes, whitespace and newlines."""
    if pd.isna(s):
        return ""
    return re.sub(r"\s+", "", str(s).replace("_x000D_", ""))


def first_id(s):
    s = str(s).replace("_x000D_", "").strip() if not pd.isna(s) else ""
    return re.split(r"[\s/]+", s)[0] if s else ""


# ---------------------------------------------------------------- 1. RevRecap
rr = pd.read_excel(REVRECAP, header=1)
rr = rr[rr["Location"].astype(str).str.startswith("QAT")].copy()
rr["AGR"] = rr["Agreement"].map(clean)
rr["INV"] = rr["IAS Invoice"].map(clean)
rr["EQ"] = rr["Equipment"].map(clean)
rr["is_att"] = rr["Product Line"].eq("CAT-ATTACH")

# ---------------------------------------------------------------- 2. Transaction History
th = pd.read_excel(TRANSHIST, header=1)
th["AGR"] = th["Trans #"].map(clean) + th["Seq#"].astype(int).astype(str).str.zfill(2)
th = th.drop_duplicates("AGR").set_index("AGR")

# ---------------------------------------------------------------- 3. Rental Data Tracker (all month sheets)
frames = []
for name, df in pd.read_excel(TRACKER, sheet_name=None, header=3).items():
    if "RENTAL AGREEMENT" not in df.columns:
        continue
    df = df.rename(columns={"ATT.1": "ATT_RATE"})
    df["sheet"] = name
    frames.append(df)
tr = pd.concat(frames, ignore_index=True)
tr["AGR"] = tr["RENTAL AGREEMENT"].map(clean).str.replace("-", "", regex=False)
tr["INV"] = tr["RENTAL INVOICE REF. NO."].map(clean)
tr["EQ"] = tr["ID"].map(first_id)
tr = tr[tr["EQ"] != ""]


def tracker_row(eq, agr, inv, single_machine):
    """Best tracker row.  Exact invoice+ID first.  A one-machine agreement then trusts
    the agreement no. (tracker IDs are sometimes swapped); a multi-machine agreement
    only has the agreement on its first tracker row, so it falls back to the ID."""
    by_agr = tr[tr.AGR == agr]
    order = [tr[(tr.INV == inv) & (tr.EQ == eq)], tr[(tr.AGR == agr) & (tr.EQ == eq)]]
    order += [by_agr, tr[tr.EQ == eq]] if single_machine else [tr[tr.EQ == eq], by_agr]
    for cand in order:
        if len(cand):
            return cand.iloc[-1], len(cand)
    return None, 0


# ---------------------------------------------------------------- 4. build machine lines, fold attachments in
lines, issues = [], []
for (agr, inv), grp in rr.groupby(["AGR", "INV"], sort=False):
    machines = grp[~grp.is_att]
    atts = grp[grp.is_att]
    recs = []
    for _, m in machines.iterrows():
        t, n = tracker_row(m.EQ, agr, inv, len(machines) == 1)
        recs.append(dict(m=m, t=t, att=None, att_rate=(t["ATT_RATE"] if t is not None else None)))
        if t is None:
            issues.append((agr, inv, m.EQ, "Machine not found in Rental Data Tracker"))
    # step 17: pair each attachment with a machine whose tracker row has an ATT rate
    takers = [r for r in recs if r["att_rate"] is not None and pd.notna(r["att_rate"]) and r["att_rate"] > 0]
    for _, a in atts.iterrows():
        target = next((r for r in takers if r["att"] is None), None)
        if target is None and len(recs) == 1:
            target = recs[0]
        if target is None:
            issues.append((agr, inv, a.EQ, f"Attachment {a.EQ} ({a.Amount}) could not be paired to a machine"))
            continue
        target["att"] = a
    for r in recs:
        m, t, a = r["m"], r["t"], r["att"]
        model = str(m.Model) if pd.notna(m.Model) else ""
        label = model + (" + H" if a is not None else "")
        agreed = float(m.Amount) + (float(a.Amount) if a is not None else 0)
        hours = float(t["CHARGEABLE HOURS"]) if t is not None and pd.notna(t["CHARGEABLE HOURS"]) else None
        opr = float(t["OPR"]) if t is not None and pd.notna(t["OPR"]) else 0
        trk_amt = float(t["INVOICE AMOUNT (QR)"]) if t is not None and pd.notna(t["INVOICE AMOUNT (QR)"]) else None
        trk_net = (trk_amt - opr * hours) if (trk_amt is not None and hours) else None
        if trk_net is not None and abs(trk_net - agreed) > 1:
            issues.append((agr, inv, m.EQ, f"RevRecap amount {agreed:,.2f} <> tracker amount less operator {trk_net:,.2f}"))
        thr = th.loc[agr] if agr in th.index else None
        if thr is None:
            issues.append((agr, inv, m.EQ, "Agreement not found in Transaction History"))
        elif clean(thr["Lead ID"]) not in grp.EQ.values:
            issues.append((agr, inv, m.EQ, f"TH Lead ID {clean(thr['Lead ID'])} not on this RevRecap invoice"))
        if t is not None and t["EQ"] != m.EQ:
            issues.append((agr, inv, m.EQ, f"Tracker ID on agreement is {t['EQ']} (RevRecap says {m.EQ})"))
        rate = FSP_RATE.get(label)
        cust = str(m.Customer)
        lines.append(dict(
            cust_no=cust[:7], cust_name=cust[7:], agr=agr,
            agr_date=thr["Agrmt Dt"] if thr is not None else None,
            invoice=int(inv.replace("PPRQ", "")), inv_date=m["Invoice Date"],
            model=label, deal="Rental", fsp=(rate * hours if rate and hours else None),
            agreed=agreed, comm=COMM_RATE.get(label), new=("Yes" if cust[:7] in NEW_CUSTOMER else "No"),
            fleet=FLEET.get(cust[:7], 1),
            se=(str(thr["Sales Rep Code"]).strip() if thr is not None else ""),
            eq=m.EQ, att_eq=(a.EQ if a is not None else ""), hours=hours, fsp_rate=rate,
            trk_rate=(t["CONTRACT PRO-RATA/HR."] if t is not None else None), opr_rate=opr,
            trk_sheet=(t["sheet"] if t is not None else ""),
        ))

# step 21: customer name, then SE
lines.sort(key=lambda d: (d["cust_name"].strip(), d["se"], d["agr"]))

# ---------------------------------------------------------------- 5. write into template
wb = openpyxl.load_workbook(TEMPLATE)
ws = wb["Monthly CommSOA"]
FIRST, TPL_LAST, TPL_TOTAL, TPL_END = 13, 42, 43, 65

# snapshot footer (total row .. end) and the data row style
footer = {}
for r in range(TPL_TOTAL, TPL_END + 1):
    footer[r] = [(c.column, c.value, copy(c._style)) for c in ws[r]]
footer_heights = {r: ws.row_dimensions[r].height for r in range(TPL_TOTAL, TPL_END + 1)}
footer_merges = [m for m in ws.merged_cells.ranges if m.min_row >= TPL_TOTAL]
row_style = {c.column: copy(c._style) for c in ws[FIRST]}
row_formulas = {c.column: c.value for c in ws[FIRST] if isinstance(c.value, str) and c.value.startswith("=")}
for m in footer_merges:
    ws.unmerge_cells(str(m))
for r in range(FIRST, ws.max_row + 1):
    for c in ws[r]:
        c.value = None
        c._style = copy(row_style.get(c.column, c._style)) if r <= TPL_LAST else c._style

n = len(lines)
last = FIRST + n - 1
total_row = last + 1
for i, d in enumerate(lines):
    r = FIRST + i
    for col, st in row_style.items():
        ws.cell(r, col)._style = copy(st)
    vals = {2: i + 1, 3: d["cust_no"], 4: d["cust_name"], 5: d["agr"], 6: d["agr_date"], 7: d["invoice"],
            8: d["inv_date"], 9: d["model"], 10: d["deal"], 11: d["fsp"], 12: d["agreed"], 15: d["comm"],
            18: d["new"], 19: d["fleet"], 21: d["se"], 22: 100}
    for col, v in vals.items():
        ws.cell(r, col).value = v
    for col, f in row_formulas.items():
        ws.cell(r, col).value = Translator(f, origin=ws.cell(FIRST, col).coordinate).translate_formula(
            ws.cell(r, col).coordinate)
    if d["fsp"] is None:
        ws.cell(r, 11).fill = YELLOW
        ws.cell(r, 11).comment = Comment(
            f"No FSP rate for '{d['model']}'. Hours = {d['hours']}", "draft")
    if d["comm"] is None:
        ws.cell(r, 15).fill = YELLOW
    # helper columns to the right of the form (outside print area)
    for col, v in zip(range(24, 32), [d["eq"], d["att_eq"], d["hours"], d["fsp_rate"], d["trk_rate"],
                                      d["opr_rate"], d["trk_sheet"], None]):
        ws.cell(r, col).value = v
for col, h in zip(range(24, 31), ["Equip. ID", "Attach. ID", "Chg. Hours", "FSP rate/hr",
                                  "Contract rate/hr", "Operator rate/hr", "Tracker sheet"]):
    ws.cell(11, col).value = h
    ws.cell(11, col).font = copy(ws.cell(11, 11).font)

# clear anything left below the new data area, then rewrite footer at new offset
for r in range(total_row, max(ws.max_row, TPL_END + n) + 1):
    for c in ws[r]:
        c.value = None
offset = total_row - TPL_TOTAL
for r, cells in footer.items():
    nr = r + offset
    for col, v, st in cells:
        c = ws.cell(nr, col)
        c._style = st
        if isinstance(v, str) and v.startswith("="):
            v = Translator(v, origin=ws.cell(r, col).coordinate).translate_formula(c.coordinate)
            # Translator leaves $-anchored rows alone; footer anchors must move with the footer
            v = re.sub(r"\$([A-Z]+)\$(\d+)", lambda mm: f"${mm[1]}${int(mm[2]) + offset}"
                       if int(mm[2]) >= TPL_TOTAL else mm[0], v)
        c.value = v
    if footer_heights[r]:
        ws.row_dimensions[nr].height = footer_heights[r]
for m in footer_merges:
    ws.merge_cells(start_row=m.min_row + offset, end_row=m.max_row + offset,
                   start_column=m.min_col, end_column=m.max_col)
ws.cell(total_row, 20).value = f"=SUM(T{FIRST}:T{last})"
ws.print_area = f"A1:T{TPL_END + offset - 4}"

month = pd.to_datetime(rr["Invoice Date"]).max()
ws["F9"] = f"ADMIN - {month:%b}".upper() + f" - {month:%Y}"
ws.oddHeader.left.text = f"  {month:%b} '{month:%y} ADMIN".upper()
ws["R8"] = pd.Timestamp.today().normalize().to_pydatetime()

# ---------------------------------------------------------------- 6. checks sheet
cs = wb.create_sheet("Checks")
cs.append(["Agreement", "Invoice", "Equipment", "Issue"])
for it in issues:
    cs.append(list(it))
cs.append([])
cs.append(["Summary"])
cs.append(["RevRecap Qatar lines", len(rr)])
cs.append(["  of which attachments", int(rr.is_att.sum())])
cs.append(["Admin machine lines", n])
cs.append(["RevRecap Qatar total", float(rr.Amount.sum())])
cs.append(["Admin Agreed SP total", sum(d["agreed"] for d in lines)])
cs.append(["Lines missing FSP rate", sum(d["fsp"] is None for d in lines)])
cs.append(["Lines missing Comm. rate", sum(d["comm"] is None for d in lines)])
for col, w in zip("ABCD", (12, 16, 12, 80)):
    cs.column_dimensions[col].width = w

wb.save(OUT)
print(f"{n} lines, {len(issues)} issues -> {OUT}")
for it in issues:
    print("  ", it)
